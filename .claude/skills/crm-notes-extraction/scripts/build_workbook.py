#!/usr/bin/env python3
"""
build_workbook.py — turns one long-format CSV into the final CRM xlsx.

Purpose
-------
This is the long/tidy-format version of the extraction hand-off,
replacing earlier JSON and wide-CSV versions. Both of those turned out
to be error-prone for a smaller local model in the same underlying
way: they require getting a multi-field structure exactly right in one
go — nested JSON keys, or a wide CSV row with many empty fields that
have to be correctly counted/aligned. This version removes that
failure mode entirely: every line names its own column explicitly, so
there is nothing to count and nothing to misalign. The model only
writes a line for a fact it actually has — never an empty placeholder.

Usage
-----
    python3 build_workbook.py <input.csv> <output.xlsx>

FORMAT of the input CSV
-------------------------
Exactly four columns, header row required:

    table,row,column,value

- table: one of "People", "Organisations", "Communities", "Projects",
  "Interactions" — exactly, case-sensitive.
- row: any whole number. Groups lines belonging to the same entity —
  e.g. all facts about the same person share one row number. Row
  numbers only need to be consistent WITHIN a table, not globally
  unique across tables (People row 1 and Organisations row 1 are
  unrelated to each other).
- column: the exact schema column name this fact belongs to (see
  ALL_HEADERS below for the authoritative list per table). An unknown
  column name causes the script to stop with an error.
- value: the actual value.

Do NOT write a line for a value you don't have — there is no way to
"leave it blank" in this format because you simply don't write the
line. Any schema column with no line for a given (table, row) is
automatically left empty AND highlighted yellow in the output. This is
the entire flagging mechanism — never write a line with an empty
value to represent "unknown"; just omit the line.

Do NOT write a line whose column is one of these forbidden columns —
the script rejects them. Some are Baserow-computed/tracked (formula
and system fields); others are out of scope for this skill for other
reasons (see FORBIDDEN_COLUMNS below and comments per table):
  - People: name-person, Last modified, Last modified by,
    Created on, Created by, edit-row-link, orgs-user-nominated,
    communities-lead-user-nominated,
    communities-participant-user-nominated, notes-user-nominated
  - Organisations: name-org, last-modified-by, last-modified-date,
    Created on, Created by, org-child-of
  - Communities: name-community, last-modified-by, last-modified-date,
    people-participants
  - Projects: name-project, last-modified-by, last-modified-date
  - Interactions: name-interaction, last-modified-by,
    last-modified-date, Created on, Created by

Example input CSV (People and Organisations only, illustrative):

table,row,column,value
People,1,name-person-first,Sara
People,1,name-person-last,Villa
People,1,notes,"Interested in training. Runs the CSCCE community."
People,1,interactions-participant,meetings with test
Organisations,1,name-org-long,CSCCE
Organisations,1,notes,Community run by Sara Villa.

(Note: no line for People's own display name — `name-person` is a
formula field computed automatically from name-person-first +
name-person-last, so it must never be written directly. Also nothing
is written for Sara's pronouns, email, affiliations, etc. — those
simply have no lines, and will show up flagged/empty in the output.
No blank-value lines needed.)

If two people are being described, the second person's lines use
row=2, e.g. "People,2,name-person-first,Andrew" — never reuse row=1
for a different entity within the same table.

Run it
------
    python3 build_workbook.py private/candidates-<name>.csv private/crm-extraction-output-<name>.xlsx

(`<name>` derived from the input notes filename — see SKILL.md Step 4
for why fixed filenames aren't used.)

The script validates every table/column name before writing anything,
and will tell you exactly which line is wrong if one doesn't match —
fix the CSV and re-run rather than trying to patch the xlsx after the
fact.

Output column order — NOT tracked against Baserow (Sept 2026 change)
----------------------------------------------------------------------
An earlier version of this script hardcoded a FULL_COLUMN_ORDER dict,
manually confirmed from Baserow's own "Hide columns" panel per table
(screenshotted directly), so a full row-range could be copy-pasted
straight from the xlsx into Baserow with every column already
aligned. That approach was dropped: the manual screenshot process
proved fragile (it drifted out of sync with the live base more than
once), and pinning down a fixed order worked against letting anyone
who templates their own Baserow instance from this one freely reorder
fields to suit themselves.

This script now simply writes each table's writable columns (per
ALL_HEADERS below) in one fixed, sensible order of its own — with no
attempt to match Baserow's live on-screen order. Practical
consequence: when copying a sheet's data into Baserow, match columns
by header name rather than pasting a full row range in one go. A
proper fix — using Baserow's own import feature, which matches
columns by field name automatically — is planned for a future
release; this is the interim approach until then.
"""

import csv
import sys
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print("ERROR: openpyxl is required (pip install openpyxl --break-system-packages)", file=sys.stderr)
    sys.exit(1)

FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
FLAG_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

# Authoritative column lists, matching references/schema.md exactly
# (rebuilt against the real Baserow field-definitions export, Sept 2026,
# refreshed again later that same month for communities-participant,
# career-stage rewording, pronouns' "self identify" option, Interactions'
# type rename/expansion, org-child-of, new Created on/by system fields,
# the writable subset of the *-user-nominated overflow fields, and the
# Communities interactions-sponsor -> interactions field-name fix).
# Forbidden columns (formula/system/out-of-scope) are deliberately NOT
# included here — the script rejects a fact line targeting any of them
# (see FORBIDDEN_COLUMNS below). This is also each table's OUTPUT
# column order in the xlsx — see the module docstring's note on why
# this no longer tries to match Baserow's own on-screen order.
ALL_HEADERS = {
    "People": [
        "name-person-first", "name-person-last", "email-1", "email-2",
        "email-3", "carrer-current-role-title", "organisations",
        "communities-contributor-profile", "communities-lead",
        "communities-participant", "interactions-lead",
        "interactions-participant", "pronouns", "pronouns-user-nominated",
        "notes", "github", "git-other-user-nominated",
        "institutional-profile", "linkedin", "socials-other-user-nominated",
        "ORCID", "projects-contact", "projects-lead", "projects-contributors",
        "career-stage", "mapping-consent",
    ],
    "Organisations": [
        "people-employees", "type", "name-org-long", "name-org-abbreviation",
        "notes", "website", "communities-sponsor", "projects-sponsor",
        "projects-funder", "interactions-sponsor", "interactions-funder",
    ],
    "Communities": [
        "description", "name-community-long", "name-community-abbreviation",
        "people-leads", "website", "organisations", "projects-sponsor",
        "interactions",
    ],
    "Projects": [
        "description", "date-start", "date-end", "status", "interactions",
        "people-contacts", "name-project-long", "name-project-abbreviation",
        "people-leads", "people-contributors", "github-issue", "notes",
        "org-sponsor", "org-funder", "communities", "website",
    ],
    "Interactions": [
        "status", "date-started", "notes", "people-leads", "projects",
        "people-participants", "name-interaction-long",
        "name-interaction-abbreviation", "type", "urls-outputs",
        "urls-internal", "date-delivered", "date-last-activity",
        "github-issue", "communities", "org-sponsor", "org-funder",
    ],
}

# Columns this skill must never write to. Not all of these are
# Baserow-computed — some (org-child-of, people-participants, and the
# *-user-nominated staging fields on People) are writable in Baserow
# but out of scope for THIS skill specifically, per schema.md's
# per-table notes.
FORBIDDEN_COLUMNS = {
    "People": {
        "name-person", "Last modified", "Last modified by",
        "Created on", "Created by", "edit-row-link",
        "orgs-user-nominated", "communities-lead-user-nominated",
        "communities-participant-user-nominated", "notes-user-nominated",
    },
    "Organisations": {
        "name-org", "last-modified-by", "last-modified-date",
        "Created on", "Created by", "org-child-of",
    },
    "Communities": {
        "name-community", "last-modified-by", "last-modified-date",
        "people-participants",
    },
    "Projects": {"name-project", "last-modified-by", "last-modified-date"},
    "Interactions": {
        "name-interaction", "last-modified-by", "last-modified-date",
        "Created on", "Created by",
    },
}

# Columns that get an explicit review flag whenever a line is written
# for them (on top of the usual "missing = flagged" mechanism) — these
# represent a genuine ambiguity worth a human glance, not just an
# absent fact. See schema.md's People notes for the exact rule each
# one follows.
EXPLICIT_FLAG_COLUMNS = {
    "People": {
        "pronouns-user-nominated", "git-other-user-nominated",
        "socials-other-user-nominated",
    },
}

REQUIRED_INPUT_COLUMNS = {"table", "row", "column", "value"}


def load_and_validate(input_path):
    """Returns (data, errors). data is {table: {row_num: {column: value}}}."""
    errors = []
    data = defaultdict(lambda: defaultdict(dict))

    with open(input_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        fieldnames = set(reader.fieldnames or [])
        missing = REQUIRED_INPUT_COLUMNS - fieldnames
        if missing:
            errors.append(
                f"Input CSV is missing required header column(s): {', '.join(sorted(missing))}. "
                f"The header row must be exactly: table,row,column,value"
            )
            return data, errors

        for line_num, line in enumerate(reader, start=2):  # start=2: header is line 1
            table = (line.get("table") or "").strip()
            row_raw = (line.get("row") or "").strip()
            column = (line.get("column") or "").strip()
            value = line.get("value")

            if table not in ALL_HEADERS:
                errors.append(
                    f"Line {line_num}: unknown table '{table}'. Valid tables: {', '.join(ALL_HEADERS)}"
                )
                continue

            if column in FORBIDDEN_COLUMNS[table]:
                errors.append(
                    f"Line {line_num}: column '{column}' on table '{table}' must not be "
                    f"written by this skill (auto-computed formula, Baserow system field, "
                    f"or explicitly out of scope — see references/schema.md). "
                    f"Remove this line."
                )
                continue

            if column not in ALL_HEADERS[table]:
                errors.append(
                    f"Line {line_num}: unknown column '{column}' for table '{table}'. "
                    f"Valid columns: {', '.join(ALL_HEADERS[table])}"
                )
                continue

            try:
                row_num = int(row_raw)
            except ValueError:
                errors.append(f"Line {line_num}: 'row' must be a whole number, got '{row_raw}'")
                continue

            if column in data[table][row_num]:
                print(
                    f"WARNING: line {line_num} duplicates table='{table}' row={row_num} "
                    f"column='{column}' — overwriting previous value "
                    f"'{data[table][row_num][column]}' with '{value}'.",
                    file=sys.stderr,
                )
            data[table][row_num][column] = value

    return data, errors


def build_workbook(data, output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for table, headers in ALL_HEADERS.items():
        rows_dict = data.get(table)
        if not rows_dict:
            continue  # no facts for this table at all — skip the sheet

        explicit_flag_cols = EXPLICIT_FLAG_COLUMNS.get(table, set())

        ws = wb.create_sheet(table)
        for col_idx, header in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col_idx, value=header)
            c.font = BOLD_FONT

        # Sort by row number for a stable, predictable sheet order
        for r_idx, row_num in enumerate(sorted(rows_dict), start=2):
            row_data = rows_dict[row_num]
            for col_idx, header in enumerate(headers, start=1):
                value = row_data.get(header)
                value = value.strip() if isinstance(value, str) else value
                is_empty = value is None or value == ""
                c = ws.cell(row=r_idx, column=col_idx, value=None if is_empty else value)
                c.font = FONT
                if is_empty:
                    c.fill = FLAG_FILL
                elif header in explicit_flag_cols:
                    # Value present but this column always gets a review
                    # flag regardless — e.g. a stated pronoun that didn't
                    # match the fixed picklist, or a non-GitHub/LinkedIn
                    # URL the model had to construct itself.
                    c.fill = FLAG_FILL

        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(18, len(header) + 4)

    wb.save(output_path)


def main():
    if len(sys.argv) != 3:
        print("Usage: python3 build_workbook.py <input.csv> <output.xlsx>", file=sys.stderr)
        sys.exit(1)

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    if not input_path.exists():
        print(f"ERROR: input file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    data, errors = load_and_validate(input_path)

    if errors:
        print("ERROR: input CSV failed validation against references/schema.md:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        print("\nFix the CSV and re-run. No file was written.", file=sys.stderr)
        sys.exit(1)

    if not data:
        print(f"ERROR: no valid data rows found in {input_path}.", file=sys.stderr)
        sys.exit(1)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(data, output_path)
    total_rows = sum(len(v) for v in data.values())
    print(f"OK: wrote {output_path} ({len(data)} sheet(s), {total_rows} row(s) total)")


if __name__ == "__main__":
    main()
</content>
